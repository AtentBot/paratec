#!/usr/bin/env python3
"""Gera o pitch da plataforma Paratec em .pptx (identidade visual da marca)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# --- Paleta da marca -------------------------------------------------------
BLUE   = RGBColor(0x1C, 0xA9, 0xE0)   # azul Paratec
BLUE_D = RGBColor(0x14, 0x7A, 0xA6)
DARK   = RGBColor(0x0E, 0x1E, 0x2E)   # ardósia
SLATE  = RGBColor(0x24, 0x37, 0x49)
GRAY   = RGBColor(0x5B, 0x66, 0x77)
LIGHT  = RGBColor(0xF3, 0xF8, 0xFC)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
AMBER  = RGBColor(0xF5, 0xA6, 0x23)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background()
    r.shadow.inherit = False
    return s

def box(s, x, y, w, h, fill=None, line=None, line_w=1.0, round_=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=6):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str): runs = [[(runs, {})]]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sp_after); p.space_before = 0
        for text, st in para:
            r = p.add_run(); r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(st.get("sz", 18))
            r.font.bold = st.get("b", False)
            r.font.italic = st.get("i", False)
            r.font.color.rgb = st.get("c", SLATE)
    return tb

def bolt(s, x, y, w, h, color=BLUE):
    """Raio estilizado (marca) via freeform."""
    fx = lambda px: Emu(int(Inches(x) + Inches(w)*px))
    fy = lambda py: Emu(int(Inches(y) + Inches(h)*py))
    pts = [(0,.45),(.55,.0),(.42,.38),(1,.30),(.45,1),(.58,.60),(0,.68)]
    fb = s.shapes.build_freeform(fx(pts[0][0]), fy(pts[0][1]))
    fb.add_line_segments([(fx(a), fy(b)) for a,b in pts[1:]], close=True)
    shp = fb.convert_to_shape()
    shp.fill.solid(); shp.fill.fore_color.rgb = color; shp.line.fill.background()
    shp.shadow.inherit = False
    return shp

def header(s, kicker, title, dark=False):
    box(s, 0, 0, 13.333, 0.28, fill=BLUE)
    txt(s, 0.7, 0.5, 12, 0.4, [[(kicker.upper(), {"sz":13,"b":True,"c":BLUE})]])
    txt(s, 0.7, 0.85, 12, 0.9, [[(title, {"sz":30,"b":True,"c":WHITE if dark else DARK})]])
    box(s, 0.72, 1.6, 0.9, 0.06, fill=AMBER)

def footer(s, n, dark=False):
    c = RGBColor(0x9A,0xA4,0xB5) if dark else GRAY
    txt(s, 0.7, 7.05, 8, 0.3, [[("Paratec · Atendimento IA no WhatsApp", {"sz":9,"c":c})]])
    txt(s, 11.5, 7.05, 1.2, 0.3, [[(f"{n:02d}", {"sz":9,"b":True,"c":BLUE})]], align=PP_ALIGN.RIGHT)

def bullets(s, x, y, w, items, sz=17, gap=10, mark="—"):
    runs = []
    for it in items:
        if isinstance(it, tuple):
            head, sub = it
            runs.append([(f"{mark}  ", {"sz":sz,"b":True,"c":BLUE}), (head, {"sz":sz,"b":True,"c":DARK})])
            runs.append([("      "+sub, {"sz":sz-3,"c":GRAY})])
        else:
            runs.append([(f"{mark}  ", {"sz":sz,"b":True,"c":BLUE}), (it, {"sz":sz,"c":SLATE})])
    txt(s, x, y, w, 5, runs, sp_after=gap)

def chip(s, x, y, w, h, title, desc, icon_color=BLUE):
    box(s, x, y, w, h, fill=WHITE, line=RGBColor(0xE2,0xE8,0xEF), line_w=1.2, round_=True)
    box(s, x, y, 0.12, h, fill=icon_color, round_=True)
    txt(s, x+0.3, y+0.18, w-0.5, 0.5, [[(title, {"sz":15,"b":True,"c":DARK})]])
    txt(s, x+0.3, y+0.62, w-0.5, h-0.7, [[(desc, {"sz":12,"c":GRAY})]])

def arrow(s, x, y, w=0.5, color=BLUE):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.35))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background(); a.shadow.inherit=False

# ===========================================================================
# 1 — CAPA
# ===========================================================================
s = slide(DARK)
box(s, 0, 0, 13.333, 7.5, fill=DARK)
bolt(s, 10.3, -1.2, 5.5, 6.0, color=RGBColor(0x17,0x2A,0x3C))   # raio decorativo grande
bolt(s, 9.7, 4.3, 4.2, 4.6, color=RGBColor(0x14,0x24,0x34))
try:
    s.shapes.add_picture("/mnt/e/Paratec/paratec-favicon-512.png", Inches(0.7), Inches(0.7),
                         height=Inches(1.5))
except Exception: pass
txt(s, 0.75, 2.7, 11, 1.4, [[("Atendimento inteligente ", {"sz":46,"b":True,"c":WHITE}),
                              ("no WhatsApp", {"sz":46,"b":True,"c":BLUE})]])
txt(s, 0.78, 4.15, 10.5, 0.8, [[("Plataforma de agentes de IA para os vendedores da Paratec —", {"sz":19,"c":RGBColor(0xC7,0xD3,0xDF)})],
                                [("cadastra clientes, consulta o catálogo e responde 24/7.", {"sz":19,"c":RGBColor(0xC7,0xD3,0xDF)})]])
box(s, 0.78, 5.5, 0.9, 0.06, fill=AMBER)
txt(s, 0.78, 5.75, 11, 0.5, [[("SPDA · Para-raios · NBR 5419", {"sz":14,"b":True,"c":BLUE})]])
txt(s, 0.78, 6.9, 11, 0.4, [[("Pitch do projeto · 2026", {"sz":12,"c":GRAY})]])

# ===========================================================================
# 2 — DESAFIO
# ===========================================================================
s = slide()
header(s, "Contexto", "O desafio do atendimento")
bullets(s, 0.9, 2.1, 11.6, [
    ("Clientes chegam pelo WhatsApp a qualquer hora", "picos fora do horário comercial ficam sem resposta"),
    ("Dúvidas repetidas de catálogo consomem o vendedor", "produtos, materiais, dimensões, códigos (SKU)"),
    ("Cadastro de clientes desestruturado", "faltam razão social, CNPJ, e-mail e contato organizados"),
    ("Pedidos, entregas e boletos sem triagem", "tudo cai no mesmo canal, sem priorização"),
], gap=14)
box(s, 0.9, 6.15, 11.5, 0.75, fill=LIGHT, line=BLUE, line_w=1.2, round_=True)
txt(s, 1.15, 6.28, 11, 0.5, [[("Oportunidade:  ", {"sz":16,"b":True,"c":BLUE}),
     ("automatizar o 1º atendimento com IA, 24/7, sem perder a qualidade nem inventar informação.", {"sz":16,"c":SLATE})]])
footer(s, 2)

# ===========================================================================
# 3 — SOLUÇÃO
# ===========================================================================
s = slide()
header(s, "A solução", "Um agente de IA que atende no WhatsApp")
cards = [
    ("Cadastra o cliente novo", "Coleta razão social, CNPJ, e-mail e contato — validando CNPJ e e-mail."),
    ("Consulta o catálogo real", "Responde com produtos, SKUs, materiais e dimensões. Nunca inventa."),
    ("Encaminha para a equipe", "Pedidos, entregas e 2ª via de boleto viram fila humana priorizada."),
    ("Tudo num painel", "Dashboard, conversas, clientes, catálogo e fila em tempo real."),
]
xs = [0.9, 7.0]; ys=[2.15, 4.35]
for i,(t,d) in enumerate(cards):
    chip(s, xs[i%2], ys[i//2], 5.45, 1.9, t, d)
box(s, 0.9, 6.55, 11.5, 0.6, fill=DARK, round_=True)
txt(s, 0.9, 6.62, 11.5, 0.45, [[("✓  Já em produção, respondendo no WhatsApp de verdade.", {"sz":15,"b":True,"c":WHITE})]], align=PP_ALIGN.CENTER)
footer(s, 3)

# ===========================================================================
# 4 — FLUXO DE ATENDIMENTO (arquitetura)
# ===========================================================================
s = slide(LIGHT)
header(s, "Como funciona", "Fluxo de atendimento ponta a ponta")
stages = [("WhatsApp","cliente"),("Evolution API","gateway"),("N8N","orquestração"),
          ("Agente IA","LangGraph + Gemini"),("Catálogo","Postgres")]
n=len(stages); x0=0.75; gap=0.28
bw=(13.333 - 2*x0 - (n-1)*(0.55+gap))/n
x=x0; y=2.7
centers=[]
for i,(t,d) in enumerate(stages):
    fill = BLUE if t=="Agente IA" else WHITE
    tcol = WHITE if t=="Agente IA" else DARK
    box(s, x, y, bw, 1.5, fill=fill, line=None if t=="Agente IA" else RGBColor(0xD6,0xE0,0xEA), line_w=1.2, round_=True)
    txt(s, x, y+0.42, bw, 0.5, [[(t, {"sz":15,"b":True,"c":tcol})]], align=PP_ALIGN.CENTER)
    txt(s, x, y+0.9, bw, 0.4, [[(d, {"sz":11,"c":(RGBColor(0xDC,0xEB,0xF5) if t=='Agente IA' else GRAY)})]], align=PP_ALIGN.CENTER)
    centers.append(x+bw)
    x += bw
    if i < n-1:
        arrow(s, x+gap*0.15, y+0.57, w=0.55); x += 0.55+gap
# resposta de volta
txt(s, x0, 4.35, 13.333-2*x0, 0.4, [[("↩  a resposta volta pelo mesmo caminho até o cliente no WhatsApp",
      {"sz":13,"b":True,"c":BLUE_D})]], align=PP_ALIGN.CENTER)
# gate de cadastro
box(s, 0.9, 5.05, 11.5, 1.55, fill=WHITE, line=BLUE, line_w=1.4, round_=True)
txt(s, 1.15, 5.2, 11, 0.4, [[("Regra de negócio: cadastro antes do catálogo", {"sz":15,"b":True,"c":BLUE})]])
g=[("Número novo","chega a mensagem"),("Cadastro","valida CNPJ/e-mail"),("Liberado","consulta o catálogo")]
gx=1.4
for i,(t,d) in enumerate(g):
    box(s, gx, 5.7, 3.0, 0.75, fill=LIGHT, round_=True)
    txt(s, gx, 5.78, 3.0, 0.35, [[(t,{"sz":13,"b":True,"c":DARK})]], align=PP_ALIGN.CENTER)
    txt(s, gx, 6.12, 3.0, 0.3, [[(d,{"sz":10,"c":GRAY})]], align=PP_ALIGN.CENTER)
    gx+=3.0
    if i<2: arrow(s, gx-0.02, 5.95, w=0.35, color=AMBER); gx+=0.35
footer(s, 4)

# ===========================================================================
# 5 — FONTES DE CONHECIMENTO
# ===========================================================================
s = slide()
header(s, "Base do agente", "Fontes de conhecimento")
box(s, 0.9, 2.1, 11.5, 1.25, fill=DARK, round_=True)
for i,(num,lab) in enumerate([("129","produtos"),("293","variantes / SKUs"),("16","categorias")]):
    cx=1.2+i*3.85
    txt(s, cx, 2.28, 3.4, 0.7, [[(num,{"sz":40,"b":True,"c":BLUE})]], align=PP_ALIGN.CENTER)
    txt(s, cx, 3.0, 3.4, 0.3, [[(lab,{"sz":13,"c":WHITE})]], align=PP_ALIGN.CENTER)
bullets(s, 0.95, 3.75, 11.6, [
    ("Catálogo de produtos (Postgres)", "extraído do site oficial; famílias, variantes/SKU, material e dimensão"),
    ("Base de clientes", "cadastros estruturados: razão social, CNPJ, e-mail, contato e status"),
    ("Histórico de conversas e eventos", "cada atendimento é persistido para métricas e continuidade"),
    ("Roadmap: base de conhecimento vetorial (RAG)", "manuais técnicos e normas (NBR 5419) para respostas ainda mais ricas"),
], sz=16, gap=9)
footer(s, 5)

# ===========================================================================
# 6 — FUNÇÕES (ferramentas)
# ===========================================================================
s = slide(LIGHT)
header(s, "Capacidades", "Funções do agente (ferramentas)")
groups = [
    ("Cadastro", BLUE, ["verificar_cliente", "cadastrar_cliente", "validação de CNPJ e e-mail"]),
    ("Catálogo", BLUE_D, ["buscar_produtos", "detalhes_produto", "buscar_por_sku",
                          "listar_categorias", "produtos_por_categoria"]),
    ("Atendimento humano", AMBER, ["registrar_pedido", "consultar_entrega", "segunda_via_boleto",
                                   "→ geram fila para a equipe"]),
]
x=0.9
for t,col,items in groups:
    box(s, x, 2.15, 3.75, 4.4, fill=WHITE, line=RGBColor(0xDD,0xE5,0xEE), line_w=1.2, round_=True)
    box(s, x, 2.15, 3.75, 0.65, fill=col, round_=True)
    txt(s, x, 2.27, 3.75, 0.45, [[(t,{"sz":16,"b":True,"c":WHITE})]], align=PP_ALIGN.CENTER)
    runs=[[("•  ",{"sz":13,"b":True,"c":col}),(it,{"sz":13,"c":SLATE,"b": it.startswith('→')})] for it in items]
    txt(s, x+0.3, 3.0, 3.2, 3.4, runs, sp_after=11)
    x+=3.9
footer(s, 6)

# ===========================================================================
# 7 — AÇÕES & CAPACIDADES
# ===========================================================================
s = slide()
header(s, "O que ele faz", "Ações & capacidades")
acts = [
    ("Atende 24/7 em português", "cordial, objetivo, no tom da marca"),
    ("Cadastro antes do catálogo", "gate que estrutura a base de clientes"),
    ("Respostas com dados reais", "nunca inventa produto, código ou preço"),
    ("Handoff inteligente", "pedidos, entrega e boletos vão para humano"),
    ("Resposta humana pelo painel", "o atendente assume e responde no WhatsApp"),
    ("Exportação de dados (CSV)", "clientes e fila, prontos para a operação"),
]
xs=[0.9, 7.0]; y=2.15
for i,(t,d) in enumerate(acts):
    chip(s, xs[i%2], y+(i//2)*1.45, 5.45, 1.25, t, d)
footer(s, 7)

# ===========================================================================
# 8 — PAINEL ADMINISTRATIVO
# ===========================================================================
s = slide(LIGHT)
header(s, "Operação", "Painel administrativo")
secs = [("Dashboard","métricas de atendimento"),("Conversas","atendimentos do WhatsApp"),
        ("Clientes","cadastros e status"),("Catálogo","produtos, variantes e categorias"),
        ("Fila humana","pedidos, entrega e boletos")]
x=0.9; w=2.28
for i,(t,d) in enumerate(secs):
    box(s, x, 2.3, w, 2.0, fill=WHITE, line=RGBColor(0xDD,0xE5,0xEE), line_w=1.2, round_=True)
    box(s, x, 2.3, w, 0.16, fill=BLUE, round_=True)
    txt(s, x, 2.75, w, 0.6, [[(t,{"sz":15,"b":True,"c":DARK})]], align=PP_ALIGN.CENTER)
    txt(s, x+0.15, 3.35, w-0.3, 0.8, [[(d,{"sz":11,"c":GRAY})]], align=PP_ALIGN.CENTER)
    x+=w+0.18
txt(s, 0.9, 4.7, 11.6, 0.4, [[("Dados reais do agente · tema claro/escuro · construído em Next.js",
     {"sz":14,"b":True,"c":BLUE_D})]])
bullets(s, 0.95, 5.2, 11.6, [
    "Acompanhe cada conversa e assuma o atendimento quando necessário",
    "Veja os clientes cadastrados e a fila de pedidos em tempo real",
], sz=15, gap=8)
footer(s, 8)

# ===========================================================================
# 9 — ARQUITETURA & STACK
# ===========================================================================
s = slide(DARK)
header(s, "Tecnologia", "Arquitetura & stack", dark=True)
rows = [
    ("Inteligência", "Google Gemini (gemini-flash-latest) + LangGraph — agente único de atendimento"),
    ("Backend", "FastAPI (Python) · PostgreSQL (catálogo + operação)"),
    ("Frontend", "Next.js + Tailwind — painel administrativo"),
    ("WhatsApp", "Evolution API (mensageria) + N8N (orquestração do fluxo)"),
    ("Infraestrutura", "Docker Swarm (Hetzner) · Traefik · Portainer — já em produção"),
]
y=2.2
for t,d in rows:
    box(s, 0.9, y, 11.55, 0.82, fill=RGBColor(0x15,0x27,0x39), round_=True)
    box(s, 0.9, y, 0.12, 0.82, fill=BLUE, round_=True)
    txt(s, 1.2, y+0.1, 3.0, 0.6, [[(t,{"sz":15,"b":True,"c":BLUE})]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 4.2, y+0.1, 8.0, 0.6, [[(d,{"sz":14,"c":RGBColor(0xD6,0xDE,0xE7)})]], anchor=MSO_ANCHOR.MIDDLE)
    y+=0.92
footer(s, 9, dark=True)

# ===========================================================================
# 10 — STATUS & ROADMAP
# ===========================================================================
s = slide()
header(s, "Onde estamos", "Status & próximos passos")
box(s, 0.9, 2.15, 5.6, 4.4, fill=LIGHT, line=RGBColor(0x1F,0xA0,0x6A), line_w=1.4, round_=True)
txt(s, 1.15, 2.3, 5.2, 0.5, [[("✓  Pronto e funcionando", {"sz":17,"b":True,"c":RGBColor(0x1F,0xA0,0x6A)})]])
bullets(s, 1.15, 2.9, 5.1, [
    "Agente em produção, atendendo no WhatsApp",
    "Cadastro de clientes + consulta ao catálogo",
    "Painel administrativo com dados reais",
    "Persistência, testes e deploy automatizado",
], sz=14, gap=9, mark="•")
box(s, 6.85, 2.15, 5.6, 4.4, fill=LIGHT, line=BLUE, line_w=1.4, round_=True)
txt(s, 7.1, 2.3, 5.2, 0.5, [[("→  Próximos passos", {"sz":17,"b":True,"c":BLUE})]])
bullets(s, 7.1, 2.9, 5.1, [
    "Número de WhatsApp dedicado da Paratec",
    "Otimizar o tempo de resposta",
    "Base de conhecimento (RAG): normas e manuais",
    "Integrações reais: ERP, logística e financeiro",
], sz=14, gap=9, mark="•")
footer(s, 10)

# ===========================================================================
# 11 — FECHAMENTO
# ===========================================================================
s = slide(DARK)
box(s, 0,0,13.333,7.5, fill=DARK)
bolt(s, 5.6, 0.9, 2.2, 2.4, color=BLUE)
txt(s, 1, 3.5, 11.333, 1.2, [[("Do primeiro “oi” ao orçamento — ", {"sz":36,"b":True,"c":WHITE}),
                               ("com IA, 24/7.", {"sz":36,"b":True,"c":BLUE})]], align=PP_ALIGN.CENTER)
txt(s, 1, 4.8, 11.333, 0.6, [[("Plataforma de atendimento inteligente da Paratec", {"sz":18,"c":RGBColor(0xC7,0xD3,0xDF)})]], align=PP_ALIGN.CENTER)
box(s, 6.0, 5.7, 1.3, 0.06, fill=AMBER)
txt(s, 1, 6.0, 11.333, 0.4, [[("SPDA · Para-raios · NBR 5419", {"sz":13,"b":True,"c":BLUE})]], align=PP_ALIGN.CENTER)

prs.save("/mnt/e/Paratec/Paratec-Pitch.pptx")
print("OK -> /mnt/e/Paratec/Paratec-Pitch.pptx  |  slides:", len(prs.slides._sldIdLst))
